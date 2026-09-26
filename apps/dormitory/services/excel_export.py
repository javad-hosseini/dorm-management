import io
import re
import urllib.parse
import jdatetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.http import HttpResponse

from apps.dormitory.models import Resident


def get_room_sort_key(resident):
    """
    Natural sort key for resident room numbers, followed by last name and first name.
    Handles numeric strings ('101', '102'), mixed ('101-B'), and None rooms gracefully.
    """
    if not resident.room or not resident.room.room_number:
        return (999999, "", resident.last_name or "", resident.first_name or "")

    rn = str(resident.room.room_number).strip()
    digits = re.findall(r'\d+', rn)
    num = int(digits[0]) if digits else 999999
    return (num, rn, resident.last_name or "", resident.first_name or "")


def get_active_debtor_residents(queryset=None):
    """
    Returns active residents currently living in the dormitory who are in debt,
    sorted ascending by room number, then resident name.
    """
    if queryset is None:
        base_qs = Resident.objects.filter(
            status=Resident.Status.ACTIVE
        ).select_related('room', 'dormitory')
    else:
        base_qs = queryset.filter(
            status=Resident.Status.ACTIVE
        ).select_related('room', 'dormitory')

    debtors = [r for r in base_qs if r.is_in_debt]
    debtors.sort(key=get_room_sort_key)
    return debtors


def get_period_debt_tomans(resident) -> int:
    """
    Calculates current period debt:
    Room monthly/period rent minus any payments made towards this period.
    Example: Rent is 3,000,000, resident paid 1,000,000 -> returns 2,000,000.
    """
    if not resident.room or not resident.room.monthly_rent:
        return 0

    rent_tomans = resident.current_monthly_rent_tomans
    if rent_tomans <= 0:
        return 0

    today = jdatetime.date.today()
    settled = resident.settled_until

    # If no settled date, fallback to entry date
    if not settled:
        settled = resident.entry_date

    # If still no date, full month rent is owed
    if not settled:
        return rent_tomans

    from apps.dormitory.jalali_utils import get_days_in_jalali_month
    import datetime

    days_in_cur_month = get_days_in_jalali_month(today.year, today.month)
    cur_month_end = jdatetime.date(today.year, today.month, days_in_cur_month) + datetime.timedelta(days=1)
    cur_month_start = jdatetime.date(today.year, today.month, 1)

    if settled >= cur_month_end:
        # Fully settled for current month
        return 0
    elif settled <= cur_month_start:
        # Settled before current month started -> full current month is owed
        # minus any approved payments made specifically in this month
        paid_in_month = sum(
            t.amount // 10
            for t in resident.transactions.filter(
                transaction_type='RENT',
                is_approved=True,
                payment_date__year=today.year,
                payment_date__month=today.month,
            )
        )
        return max(0, rent_tomans - paid_in_month)
    else:
        # settled is mid-month (e.g. 1405/07/11 after a 10-day partial payment)
        remaining_days = (cur_month_end - settled).days
        daily_rate = rent_tomans / days_in_cur_month
        remaining_tomans = int(round(remaining_days * daily_rate))
        return max(0, min(rent_tomans, remaining_tomans))


def generate_debtors_excel_workbook(queryset=None) -> openpyxl.Workbook:
    """
    Builds an openpyxl Workbook containing all active debtors formatted with:
    - Right-to-Left (RTL) layout
    - 8 Columns: شماره اتاق, اسامی بدهکارا, تاریخی که تسویه شدن, بدهی دوره جاری (تومان),
      کل بدهی معوقه (تومان), شماره تلفن همراه, شماره تلفن ضروری, توضیحات
    - Cells with amounts formatted with 'تومان' unit (#,##0 "تومان")
    - Styled executive headers, number formatting, borders, and total summary row.
    """
    debtors = get_active_debtor_residents(queryset)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "لیست بدهکاران"

    # Set Right-To-Left (RTL) layout for Persian Excel view
    ws.sheet_view.rightToLeft = True

    # Styles definition
    font_family = "Calibri"
    header_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")  # Deep Navy
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    border_thin = Side(style="thin", color="CBD5E1")
    cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

    row_font = Font(name=font_family, size=11, color="0F172A")
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    toman_fmt = '#,##0 "تومان"'

    headers = [
        "شماره اتاق",
        "اسامی بدهکارا",
        "تاریخی که تسویه شدن",
        "بدهی دوره جاری (تومان)",
        "کل بدهی معوقه (تومان)",
        "شماره تلفن همراه",
        "شماره تلفن ضروری",
        "توضیحات",
    ]

    # Write Header Row
    ws.row_dimensions[1].height = 32
    for col_idx, header_title in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header_title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = cell_border

    # Write Data Rows
    current_row = 2
    for r in debtors:
        ws.row_dimensions[current_row].height = 24
        fill = alt_fill if current_row % 2 == 0 else white_fill

        # 1. Room Number
        room_str = str(r.room.room_number) if (r.room and r.room.room_number) else "بدون اتاق"

        # 2. Debtor Name
        name_str = r.full_name

        # 3. Settled Until Date
        if r.settled_until:
            settled_str = r.settled_until.strftime('%Y/%m/%d')
        elif r.entry_date:
            settled_str = f"{r.entry_date.strftime('%Y/%m/%d')} (ورود)"
        else:
            settled_str = "ثبت نشده"

        # 4. Period Debt (Tomans) - بدهی کل دوره با کسر پرداختی
        period_debt_tomans = get_period_debt_tomans(r)

        # 5. Total Accumulated Overdue Debt (Tomans) - مبلغ بدهی
        debt_tomans = r.total_debt_amount_tomans

        # 6. Phone Number
        phone_str = str(r.phone_number or "-")

        # 7. Emergency Phone Number
        emergency_str = str(r.parent_phone_number or "-")

        # 8. Remarks / Notes (blank for manual notes)
        notes_str = ""

        row_data = [
            (room_str, Alignment(horizontal="center", vertical="center"), "@"),
            (name_str, Alignment(horizontal="right", vertical="center"), None),
            (settled_str, Alignment(horizontal="center", vertical="center"), None),
            (period_debt_tomans, Alignment(horizontal="right", vertical="center", readingOrder=2), toman_fmt),
            (debt_tomans, Alignment(horizontal="right", vertical="center", readingOrder=2), toman_fmt),
            (phone_str, Alignment(horizontal="center", vertical="center"), "@"),
            (emergency_str, Alignment(horizontal="center", vertical="center"), "@"),
            (notes_str, Alignment(horizontal="right", vertical="center"), None),
        ]

        for col_idx, (val, align, num_fmt) in enumerate(row_data, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = row_font
            cell.fill = fill
            cell.alignment = align
            cell.border = cell_border
            if num_fmt:
                cell.number_format = num_fmt

        current_row += 1

    # Write Summary / Total Row
    total_row = current_row
    ws.row_dimensions[total_row].height = 28

    total_label = f"مجموع کل ({len(debtors)} نفر بدهکار)"
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=3)

    summary_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    summary_font = Font(name=font_family, size=11, bold=True, color="0F172A")
    summary_border_top = Side(style="thin", color="0F172A")
    summary_border_bottom = Side(style="double", color="0F172A")
    summary_border = Border(
        left=border_thin,
        right=border_thin,
        top=summary_border_top,
        bottom=summary_border_bottom
    )

    # Merged label in columns 1-3
    for c in range(1, 4):
        cell = ws.cell(row=total_row, column=c)
        cell.fill = summary_fill
        cell.border = summary_border
        if c == 1:
            cell.value = total_label
            cell.font = summary_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

    # Column 4: Period Debt SUM formula
    period_debt_cell = ws.cell(row=total_row, column=4)
    period_debt_cell.value = f"=SUM(D2:D{total_row - 1})" if len(debtors) > 0 else 0
    period_debt_cell.font = Font(name=font_family, size=11, bold=True, color="B91C1C")
    period_debt_cell.fill = summary_fill
    period_debt_cell.alignment = Alignment(horizontal="right", vertical="center", readingOrder=2)
    period_debt_cell.border = summary_border
    period_debt_cell.number_format = toman_fmt

    # Column 5: Total Overdue Debt SUM formula
    total_debt_cell = ws.cell(row=total_row, column=5)
    total_debt_cell.value = f"=SUM(E2:E{total_row - 1})" if len(debtors) > 0 else 0
    total_debt_cell.font = Font(name=font_family, size=11, bold=True, color="B91C1C")
    total_debt_cell.fill = summary_fill
    total_debt_cell.alignment = Alignment(horizontal="right", vertical="center", readingOrder=2)
    total_debt_cell.border = summary_border
    total_debt_cell.number_format = toman_fmt

    # Columns 6, 7, 8 in summary row
    for c in range(6, 9):
        cell = ws.cell(row=total_row, column=c, value="")
        cell.fill = summary_fill
        cell.border = summary_border

    # Set Column Widths
    col_widths = {
        1: 16,  # شماره اتاق
        2: 28,  # اسامی بدهکارا
        3: 24,  # تاریخی که تسویه شدن
        4: 26,  # بدهی دوره جاری (تومان)
        5: 26,  # کل بدهی معوقه (تومان)
        6: 18,  # شماره تلفن همراه
        7: 18,  # شماره تلفن ضروری
        8: 35,  # توضیحات
    }
    for col_idx, width in col_widths.items():
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    return wb



def export_debtors_to_excel_response(queryset=None) -> HttpResponse:
    """
    Generates the Excel file for debtors and returns it as a downloadable HttpResponse.
    The filename includes today's Jalali date (e.g. بدهکاران_1405_07_05.xlsx).
    """
    wb = generate_debtors_excel_workbook(queryset)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    today = jdatetime.date.today()
    jalali_str = f"{today.year:04d}_{today.month:02d}_{today.day:02d}"
    filename_fa = f"بدهکاران_{jalali_str}.xlsx"
    filename_ascii = f"debtors_{jalali_str}.xlsx"

    quoted_filename = urllib.parse.quote(filename_fa)

    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response['Content-Disposition'] = (
        f'attachment; filename="{filename_ascii}"; filename*=UTF-8\'\'{quoted_filename}'
    )
    return response
